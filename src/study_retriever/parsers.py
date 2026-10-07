from __future__ import annotations

import csv
import io
import json
import re
import shutil
from pathlib import Path
from typing import Any, cast

from .config import SUPPORTED_SUFFIXES, AppConfig
from .models import Block, ParsedDocument


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    lines = [line.rstrip() for line in text.splitlines()]
    out: list[str] = []
    blank = False
    for line in lines:
        if line.strip():
            out.append(line)
            blank = False
        elif not blank:
            out.append("")
            blank = True
    return "\n".join(out).strip()


def _decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


_MD_IMAGE = re.compile(r"!\[([^\]]*)\]\([^)]*\)")
_GENERIC_ALT = re.compile(r"(?i)\s*(page \d+ image|image|img|figure|picture|photo)?\s*")


def _strip_markdown_images(text: str) -> str:
    """Keep meaningful alt text; drop bare image links that only dilute chunk embeddings."""
    return _MD_IMAGE.sub(lambda m: "" if _GENERIC_ALT.fullmatch(m.group(1)) else m.group(1), text)


def _limits() -> AppConfig:
    return AppConfig()


def _check_text_size(path: Path) -> None:
    if path.stat().st_size > _limits().max_text_bytes:
        raise ValueError(f"{path.name}: larger than the {_limits().max_text_bytes // 2**20} MB text-parser limit")


def _check_zip(path: Path) -> None:
    """Reject zip bombs before python-docx/pptx/openpyxl load every part into memory."""
    import zipfile

    with zipfile.ZipFile(path) as zf:
        total = sum(info.file_size for info in zf.infolist())
    if total > _limits().max_office_uncompressed_bytes:
        raise ValueError(f"{path.name}: archive expands to {total // 2**20} MB, over the safety limit")


def _generic_text(path: Path) -> ParsedDocument:
    _check_text_size(path)
    text = _decode_bytes(path.read_bytes())
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".qmd"}:
        text = _strip_markdown_images(text)
    if suffix == ".csv":
        try:
            rows = list(csv.reader(io.StringIO(text)))
            text = "\n".join("\t".join(cell for cell in row) for row in rows)
        except csv.Error:
            pass
    return ParsedDocument(
        title=path.stem,
        blocks=[Block(_clean_text(text), locator={"line_start": 1})],
        metadata={"format": suffix.lstrip(".") or "text"},
    )


MAX_OCR_PIXELS = 6000


def _ocr_dpi(page: Any) -> int:
    """OCR resolution that keeps the raster under MAX_OCR_PIXELS per side; 0 means skip (crafted giant page)."""
    longest = max(float(page.rect.width), float(page.rect.height), 1.0)
    dpi = min(150, int(MAX_OCR_PIXELS * 72 / longest))
    return dpi if dpi >= 40 else 0


def _tessdata() -> str | None:
    exe = shutil.which("tesseract")
    folder = Path(exe).parent / "tessdata" if exe else None
    return str(folder) if folder and folder.is_dir() else None


def _pdf(path: Path) -> ParsedDocument:
    import fitz

    doc: Any = fitz.open(path)
    try:
        title = (doc.metadata or {}).get("title") or path.stem
        toc = doc.get_toc(simple=True) or []
        headings_by_page: dict[int, str] = {}
        current = ""
        toc_index = 0
        sorted_toc = sorted((int(page), str(name).strip()) for _level, name, page in toc if page)
        for page_num in range(1, doc.page_count + 1):
            while toc_index < len(sorted_toc) and sorted_toc[toc_index][0] <= page_num:
                current = sorted_toc[toc_index][1]
                toc_index += 1
            if current:
                headings_by_page[page_num] = current

        blocks: list[Block] = []
        ocr_available = shutil.which("tesseract") is not None
        ocr_pages: list[int] = []
        max_ocr = _limits().max_ocr_pages
        for i, page in enumerate(doc, start=1):
            text = _clean_text(page.get_text("text", sort=True))
            ocr_dpi = _ocr_dpi(page) if len(text) < 20 and ocr_available and len(ocr_pages) < max_ocr else 0
            if ocr_dpi:
                try:
                    tp = page.get_textpage_ocr(language="eng", dpi=ocr_dpi, full=True, tessdata=_tessdata())
                    ocr_text = _clean_text(page.get_text("text", textpage=tp, sort=True))
                    if len(ocr_text) > len(text):
                        text = ocr_text
                        ocr_pages.append(i)
                except Exception:
                    pass
            if text:
                blocks.append(Block(text, section=headings_by_page.get(i, ""), locator={"page": i}))
        return ParsedDocument(
            title=title,
            blocks=blocks,
            metadata={"format": "pdf", "pages": doc.page_count, "ocr_pages": ocr_pages},
        )
    finally:
        doc.close()


def _legible(text: str) -> bool:
    """Reject OCR noise from plots/photos: most tokens must look like real words."""
    tokens = [t.strip(".,;:()[]{}\"'!?-_") for t in text.split()]
    tokens = [t for t in tokens if len(t) >= 2]
    good = [t for t in tokens if re.fullmatch(r"[A-Za-z][a-z]{2,}|[A-Z]{2,}|[A-Za-z]*\d+[A-Za-z\d.]*|[a-z]{2}", t)]
    words = [t for t in good if t.isalpha() and len(t) >= 3]
    return len(words) >= 5 and len(good) >= 0.7 * len(tokens)


def _image(path: Path) -> ParsedDocument:
    """OCR a PNG/JPEG through PyMuPDF + Tesseract; images without legible text yield no blocks."""
    import fitz

    blocks: list[Block] = []
    if path.stat().st_size >= 2048 and shutil.which("tesseract") is not None:
        doc: Any = fitz.open(path)
        try:
            page = doc[0]
            ocr_dpi = _ocr_dpi(page)
            if ocr_dpi:
                tp = page.get_textpage_ocr(language="eng", dpi=ocr_dpi, full=True, tessdata=_tessdata())
                text = _clean_text(str(page.get_text("text", textpage=tp, sort=True)))
                if _legible(text):
                    blocks.append(Block(text, locator={"image": path.name}))
        finally:
            doc.close()
    return ParsedDocument(title=path.stem, blocks=blocks, metadata={"format": "image-ocr"})


def _svg(path: Path) -> ParsedDocument:
    """Index visible labels, titles and descriptions of vector figures; path-only SVGs yield no blocks."""
    import xml.etree.ElementTree as ET

    texts: list[str] = []
    _check_text_size(path)
    raw = _decode_bytes(path.read_bytes())
    if "<!ENTITY" in raw:
        raise ValueError(f"{path.name}: SVG with entity declarations is refused (expansion attack)")
    try:
        for el in ET.fromstring(raw.encode("utf-8")).iter():
            if el.tag.rsplit("}", 1)[-1] in {"title", "desc", "text"}:
                value = _clean_text(" ".join("".join(el.itertext()).split()))
                if value and value not in texts:
                    texts.append(value)
    except ET.ParseError:
        pass
    blocks = [Block("\n".join(texts), locator={"figure": path.name})] if texts else []
    return ParsedDocument(title=path.stem, blocks=blocks, metadata={"format": "svg"})


def _npz(path: Path) -> ParsedDocument:
    """Describe a NumPy archive (array names, shapes, dtypes, short string samples); numbers are not indexed."""
    import numpy as np

    lines: list[str] = []
    with np.load(path, allow_pickle=False) as archive:
        for name in archive.files:
            arr = archive[name]
            line = f"{name}: shape {tuple(arr.shape)}, dtype {arr.dtype}"
            if arr.dtype.kind in "US" and arr.size:
                sample = sorted({str(v) for v in arr.ravel()[:2000]})[:25]
                line += ", values: " + ", ".join(sample)
            lines.append(line)
    blocks = [Block("NumPy archive " + path.name + "\n" + "\n".join(lines), locator={"archive": path.name})]
    return ParsedDocument(title=path.stem, blocks=blocks, metadata={"format": "npz"})


def _docx(path: Path) -> ParsedDocument:
    from docx import Document

    _check_zip(path)
    doc = Document(str(path))
    title = path.stem
    current_heading = ""
    blocks: list[Block] = []
    paragraph_index = 0
    for para in doc.paragraphs:
        paragraph_index += 1
        text = _clean_text(para.text)
        if not text:
            continue
        style = (para.style.name if para.style else "") or ""
        if style.lower().startswith("title") and title == path.stem:
            title = text
        if style.lower().startswith("heading"):
            current_heading = text
            continue
        blocks.append(Block(text, section=current_heading, locator={"paragraph": paragraph_index}))

    for table_index, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append("\t".join(_clean_text(cell.text).replace("\n", " ") for cell in row.cells))
        text = _clean_text("\n".join(rows))
        if text:
            blocks.append(Block(text, section=current_heading, locator={"table": table_index}))
    return ParsedDocument(title=title, blocks=blocks, metadata={"format": "docx"})


def _pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation

    _check_zip(path)
    prs = Presentation(str(path))
    title = path.stem
    blocks: list[Block] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        slide_title = ""
        if slide.shapes.title is not None:
            slide_title = _clean_text(slide.shapes.title.text)
            if slide_index == 1 and slide_title:
                title = slide_title
        for shape in cast(Any, slide.shapes):
            if getattr(shape, "has_text_frame", False):
                text = _clean_text(shape.text)
                if text and text != slide_title:
                    texts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    texts.append("\t".join(_clean_text(c.text).replace("\n", " ") for c in row.cells))
        try:
            notes_text: list[str] = []
            notes_slide = slide.notes_slide
            for shape in cast(Any, notes_slide.shapes):
                if getattr(shape, "has_text_frame", False):
                    t = _clean_text(shape.text)
                    if t and not t.lower().startswith("click to edit"):
                        notes_text.append(t)
            if notes_text:
                texts.append("Speaker notes:\n" + "\n".join(notes_text))
        except Exception:
            pass
        text = _clean_text("\n\n".join(texts))
        if slide_title:
            text = _clean_text(slide_title + ("\n\n" + text if text else ""))
        if text:
            blocks.append(Block(text, section=slide_title, locator={"slide": slide_index}))
    return ParsedDocument(title=title, blocks=blocks, metadata={"format": "pptx", "slides": len(prs.slides)})


def _xlsx(path: Path) -> ParsedDocument:
    from openpyxl import load_workbook

    _check_zip(path)
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        blocks: list[Block] = []
        for ws in wb.worksheets:
            batch: list[str] = []
            start_row = 1
            last_row = 0
            for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    if not batch:
                        start_row = row_index
                    batch.append("\t".join(values).rstrip())
                    last_row = row_index
                if len(batch) >= 50:
                    blocks.append(Block(_clean_text("\n".join(batch)), section=ws.title,
                                        locator={"sheet": ws.title, "row_start": start_row, "row_end": last_row}))
                    batch = []
            if batch:
                blocks.append(Block(_clean_text("\n".join(batch)), section=ws.title,
                                    locator={"sheet": ws.title, "row_start": start_row, "row_end": last_row}))
        return ParsedDocument(title=path.stem, blocks=blocks,
                              metadata={"format": path.suffix.lower().lstrip("."), "sheets": wb.sheetnames})
    finally:
        wb.close()


def _ipynb(path: Path) -> ParsedDocument:
    obj = json.loads(path.read_text(encoding="utf-8"))
    blocks: list[Block] = []
    current_heading = ""
    for i, cell in enumerate(obj.get("cells", []), start=1):
        source = _clean_text("".join(cell.get("source", [])))
        if not source:
            continue
        cell_type = cell.get("cell_type", "")
        if cell_type == "markdown":
            for line in source.splitlines():
                if line.lstrip().startswith("#"):
                    current_heading = line.lstrip("#").strip()
                    break
        prefix = f"[{cell_type} cell]\n" if cell_type else ""
        blocks.append(Block(prefix + source, section=current_heading, locator={"cell": i, "cell_type": cell_type}))
    return ParsedDocument(title=path.stem, blocks=blocks, metadata={"format": "ipynb"})


def _html(path: Path) -> ParsedDocument:
    from bs4 import BeautifulSoup

    _check_text_size(path)
    soup = BeautifulSoup(_decode_bytes(path.read_bytes()), "html.parser")
    for tag in soup(["script", "style", "noscript", "svg"]):
        tag.decompose()
    title = _clean_text(soup.title.get_text(" ", strip=True)) if soup.title else path.stem
    blocks: list[Block] = []
    current_heading = ""
    buffer: list[str] = []
    for node in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "code", "table"]):
        text = _clean_text(node.get_text("\n" if node.name in {"pre", "code", "table"} else " ", strip=True))
        if not text:
            continue
        if node.name.startswith("h"):
            if buffer:
                blocks.append(Block(_clean_text("\n\n".join(buffer)), section=current_heading))
                buffer = []
            current_heading = text
        else:
            buffer.append(text)
    if buffer:
        blocks.append(Block(_clean_text("\n\n".join(buffer)), section=current_heading))
    if not blocks:
        text = _clean_text(soup.get_text("\n", strip=True))
        if text:
            blocks.append(Block(text))
    return ParsedDocument(title=title or path.stem, blocks=blocks, metadata={"format": "html"})


PARSERS = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".pptx": _pptx,
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".ipynb": _ipynb,
    ".html": _html,
    ".htm": _html,
    ".png": _image,
    ".jpg": _image,
    ".jpeg": _image,
    ".bmp": _image,
    ".gif": _image,
    ".tif": _image,
    ".tiff": _image,
    ".svg": _svg,
    ".npz": _npz,
}


def parse_document(path: Path) -> ParsedDocument:
    path = path.expanduser().resolve()
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported file type: {suffix or '<none>'}")
    parser = PARSERS.get(suffix, _generic_text)
    doc = parser(path)
    doc.blocks = [Block(_clean_text(b.text), b.section, dict(b.locator)) for b in doc.blocks if _clean_text(b.text)]
    return doc


def is_supported(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
