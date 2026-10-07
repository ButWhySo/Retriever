# pyright: reportArgumentType=false, reportAttributeAccessIssue=false, reportOptionalMemberAccess=false
from __future__ import annotations

import json
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from study_retriever.config import AppConfig, Paths, save_config
from study_retriever.runtime import Runtime


def make_real_corpus(root: Path) -> None:
    root.mkdir(parents=True, exist_ok=True)
    (root / "attention.md").write_text(
        "# Transformer Attention\n\nScaled dot-product attention divides QK^T by sqrt(d_k) before softmax.\n"
        "The scaling prevents large dot products from saturating softmax.\n",
        encoding="utf-8",
    )

    doc = Document()
    doc.add_heading("Virtual Memory", 0)
    doc.add_heading("TLB", level=1)
    doc.add_paragraph("A translation lookaside buffer caches recent virtual-to-physical page translations.")
    doc.save(root / "virtual-memory.docx")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Kruskal MST"
    slide.placeholders[1].text = "Kruskal sorts edges by increasing weight and uses union-find to avoid cycles."
    prs.save(root / "graphs.pptx")

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Gradient descent updates parameters opposite the gradient using a step size.")
    pdf.save(root / "optimization.pdf")
    pdf.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "MLE"
    ws.append(["Concept", "Definition"])
    ws.append(["Maximum likelihood estimation", "Choose parameters maximizing observed-data likelihood"])
    wb.save(root / "statistics.xlsx")

    notebook = {
        "cells": [
            {"cell_type": "markdown", "metadata": {}, "source": ["# Recurrent Neural Networks\n", "Hidden states carry sequence context."]},
            {"cell_type": "code", "metadata": {}, "execution_count": None, "outputs": [], "source": ["h_t = tanh(W_hh @ h_prev + W_xh @ x_t)"]},
        ],
        "metadata": {},
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    (root / "rnn.ipynb").write_text(json.dumps(notebook), encoding="utf-8")

    (root / "minimax.html").write_text(
        "<html><head><title>Minimax</title></head><body><h1>Saddle Point</h1>"
        "<p>A saddle point satisfies the minimax equilibrium inequalities.</p></body></html>",
        encoding="utf-8",
    )


def build_runtime(tmp_path: Path) -> Runtime:
    home = tmp_path / "home"
    paths = Paths(home)
    cfg = AppConfig(vector_backend="hashing", watcher_enabled=False, chunk_chars=700, overlap_chars=80)
    save_config(paths, cfg)
    return Runtime(home=home, start_watcher=False)


def test_full_incremental_hybrid_pipeline(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    make_real_corpus(corpus)
    rt = build_runtime(tmp_path)
    try:
        added = rt.add_root(str(corpus), sync_now=True)
        assert added["sync"]["errors"] == 0
        status = rt.status()
        assert status["sources"] == 7
        assert status["chunks"] >= 7
        assert status["vectors"] == status["chunks"]
        assert status["vector_count_matches_chunks"] is True

        hits = rt.search.hybrid_search("scaled dot product attention sqrt dk", top_k=5)
        assert hits
        assert Path(hits[0].path).name == "attention.md"
        assert "sqrt(d_k)" in hits[0].text

        vm = rt.search.hybrid_search("translation lookaside buffer", top_k=5)
        assert any(Path(h.path).name == "virtual-memory.docx" for h in vm)

        graphs = rt.search.hybrid_search("Kruskal union-find cycles", top_k=5)
        assert any(Path(h.path).name == "graphs.pptx" and h.locator.get("slide") == 1 for h in graphs)

        bundle = rt.search.topic_bundle("gradient descent step size", max_chunks=8, max_chars=10_000)
        assert bundle["chunk_count"] >= 1
        assert "optimization.pdf" in bundle["content"]

        fetched = rt.search.fetch(hits[0].chunk_id)
        assert fetched is not None
        assert fetched["url"].startswith("file://")
        assert "Scaled dot-product attention" in fetched["content"]

        exact = rt.search.read_source(path=str(corpus / "graphs.pptx"), slide=1)
        assert "Kruskal" in exact["content"]
        assert exact["metadata"]["slides"] == 1

        line_exact = rt.search.read_source(path=str(corpus / "attention.md"), line_start=2, line_end=4)
        assert "Scaled dot-product attention" in line_exact["content"]
        assert "saturating softmax" in line_exact["content"]

        rendered = rt.search.render_pdf_page(path=str(corpus / "optimization.pdf"), page=1, dpi=96)
        assert rendered.startswith(b"\x89PNG\r\n\x1a\n")
        assert len(rendered) > 1000

        attention = corpus / "attention.md"
        attention.write_text(attention.read_text(encoding="utf-8") + "\nFlash attention reduces materialized attention-memory traffic.\n", encoding="utf-8")
        sync = rt.indexer.sync_root(corpus)
        assert sync.indexed == 1
        assert sync.errors == 0
        new_hits = rt.search.hybrid_search("materialized attention memory traffic", top_k=5)
        assert new_hits and Path(new_hits[0].path).name == "attention.md"

        (corpus / "minimax.html").unlink()
        sync = rt.indexer.sync_root(corpus)
        assert sync.deleted == 1
        assert rt.status()["sources"] == 6
    finally:
        rt.close()


def test_unchanged_files_do_not_reindex(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    (corpus / "notes.txt").write_text("Bayes theorem relates posterior, likelihood, prior, and evidence.", encoding="utf-8")
    rt = build_runtime(tmp_path)
    try:
        first = rt.add_root(str(corpus), sync_now=True)["sync"]
        assert first["indexed"] == 1
        second = rt.indexer.sync_root(corpus)
        assert second.indexed == 0
        assert second.unchanged == 1
    finally:
        rt.close()
