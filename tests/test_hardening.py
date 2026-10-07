from __future__ import annotations

from pathlib import Path

from docx import Document

from study_retriever.api_models import IndexStats
from study_retriever.cli import _run_live_self_test
from study_retriever.config import AppConfig, Paths, save_config
from study_retriever.runtime import Runtime


def test_default_embedding_resources_are_bounded() -> None:
    assert AppConfig().embedding_batch_size == 16


def make_runtime(tmp_path: Path) -> Runtime:
    home = tmp_path / "home"
    paths = Paths(home)
    save_config(paths, AppConfig(vector_backend="hashing", watcher_enabled=False, chunk_chars=500, overlap_chars=60))
    return Runtime(home=home, start_watcher=False)


def test_filters_special_queries_and_precise_citations(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    (corpus / "nlp").mkdir(parents=True)
    (corpus / "os").mkdir()
    (corpus / "nlp" / "attention.md").write_text(
        "# Attention\nQK^T / sqrt(d_k) is scaled before softmax.\nSecond line has multi-head attention.\n",
        encoding="utf-8",
    )
    (corpus / "os" / "attention.txt").write_text(
        "The operating system attention test file discusses interrupts, not transformers.\n",
        encoding="utf-8",
    )
    doc = Document()
    doc.add_heading("Encoder Decoder", 0)
    doc.add_heading("Cross Attention", 1)
    doc.add_paragraph("Decoder queries attend to encoder keys and values.")
    doc.save(corpus / "nlp" / "cross-attention.docx")

    rt = make_runtime(tmp_path)
    try:
        result = rt.add_root(str(corpus), sync_now=True)
        assert result["sync"]["errors"] == 0

        punctuated = rt.search.hybrid_search("QK^T sqrt(d_k) softmax", top_k=5)
        assert punctuated and Path(punctuated[0].path).name == "attention.md"
        assert "lines 1-" in punctuated[0].citation()

        filtered = rt.search.hybrid_search(
            "attention",
            top_k=10,
            path_prefix=str(corpus / "nlp"),
            file_types=["md", ".docx"],
        )
        assert filtered
        assert all(Path(h.path).is_relative_to(corpus / "nlp") for h in filtered)
        assert all(Path(h.path).suffix in {".md", ".docx"} for h in filtered)

        cross = rt.search.hybrid_search("decoder queries encoder keys values", top_k=5)
        doc_hit = next(h for h in cross if Path(h.path).name == "cross-attention.docx")
        assert doc_hit.locator["paragraph"] == 3
        assert "paragraph 3" in doc_hit.citation()
        assert "Cross Attention" in doc_hit.citation()
    finally:
        rt.close()


def test_remove_root_never_deletes_original_and_config_recovers_roots(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    note = corpus / "proofs.txt"
    note.write_text("An exchange argument proves optimality by swapping choices without worsening cost.", encoding="utf-8")
    home = tmp_path / "home"

    rt = make_runtime(tmp_path)
    try:
        rt.add_root(str(corpus), sync_now=True)
        assert rt.status()["sources"] == 1
    finally:
        rt.close()

    # Simulate catalog loss while retaining durable config; startup must restore configured roots.
    db = home / "catalog.sqlite3"
    for suffix in ("", "-wal", "-shm"):
        candidate = Path(str(db) + suffix)
        if candidate.exists():
            candidate.unlink()

    rt = Runtime(home=home, start_watcher=False)
    try:
        assert str(corpus.resolve()) in rt.status()["roots"]
        synced = rt.indexer.sync_all()
        assert synced["total"]["indexed"] == 1
        removed = rt.remove_root(str(corpus), delete_indexed=True)
        assert removed["deleted_sources"] == 1
        assert note.exists()
        assert "exchange argument" in note.read_text(encoding="utf-8")
    finally:
        rt.close()


def test_live_self_test_and_typed_status_on_real_backend(tmp_path: Path) -> None:
    rt = make_runtime(tmp_path)
    try:
        result = _run_live_self_test(rt)
        assert result["dense_rank"] == 1
        assert result["lexical_rank"] == 1
        from study_retriever import __version__
        status = IndexStats.model_validate(rt.status())
        assert status.version == __version__
        assert status.vector_count_matches_chunks is True
        assert status.sources == 0
        assert status.chunks == 0
        assert status.error_sources == []
    finally:
        rt.close()


def test_directory_watcher_resync_handles_nested_create_and_delete(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    nested = corpus / "old"
    nested.mkdir(parents=True)
    (nested / "a.md").write_text("original watcher source about paging", encoding="utf-8")
    rt = make_runtime(tmp_path)
    try:
        rt.add_root(str(corpus), sync_now=True)
        assert rt.status()["sources"] == 1

        new_dir = corpus / "new"
        new_dir.mkdir()
        (new_dir / "b.md").write_text("new directory event source about attention", encoding="utf-8")
        rt.watcher._sync_directory_event(new_dir)
        assert rt.status()["sources"] == 2

        import shutil
        shutil.rmtree(nested)
        rt.watcher._sync_directory_event(nested)
        assert rt.status()["sources"] == 1
        stale = rt.search.hybrid_search("original watcher paging", top_k=5)
        assert all(Path(h.path).name != "a.md" for h in stale)
    finally:
        rt.close()


def test_fetch_contains_precise_citation(tmp_path: Path) -> None:
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    note = corpus / "attention.md"
    note.write_text("# Attention\nScaled dot product attention uses query key value vectors.\n", encoding="utf-8")
    rt = make_runtime(tmp_path)
    try:
        rt.add_root(str(corpus), sync_now=True)
        hit = rt.search.hybrid_search("query key value", top_k=3)[0]
        fetched = rt.search.fetch(hit.chunk_id)
        assert fetched is not None
        assert fetched["citation"]
        assert str(note.resolve()) in fetched["citation"]
    finally:
        rt.close()


def test_two_runtime_instances_see_each_others_index_updates(tmp_path: Path) -> None:
    home = tmp_path / "shared-home"
    paths = Paths(home)
    save_config(paths, AppConfig(vector_backend="hashing", watcher_enabled=False, chunk_chars=500, overlap_chars=60))
    corpus = tmp_path / "corpus"
    corpus.mkdir()
    note = corpus / "shared.md"
    note.write_text("cross process retrieval sees galactic semaphore token", encoding="utf-8")

    reader = Runtime(home=home, start_watcher=False)
    writer = Runtime(home=home, start_watcher=False)
    try:
        writer.add_root(str(corpus), sync_now=True)
        hits = reader.search.hybrid_search("galactic semaphore", top_k=5)
        assert any(Path(h.path) == note.resolve() for h in hits)

        note.write_text("cross process retrieval now sees cobalt lighthouse token", encoding="utf-8")
        writer.indexer.sync_root(corpus)
        hits = reader.search.hybrid_search("cobalt lighthouse", top_k=5)
        assert any(Path(h.path) == note.resolve() for h in hits)
    finally:
        writer.close()
        reader.close()


def _concurrent_sync_worker(home: str, root: str, gate, out) -> None:
    """Real child-process index writer used by concurrency integration test."""
    runtime = Runtime(home=Path(home), start_watcher=False)
    try:
        gate.wait(timeout=10)
        result = runtime.indexer.sync_root(Path(root), force=True)
        out.put({"indexed": result.indexed, "errors": result.errors})
    except Exception as exc:  # reported to parent; not a fake success path
        out.put({"error": f"{type(exc).__name__}: {exc}"})
    finally:
        runtime.close()


def test_concurrent_process_sync_keeps_catalog_and_vectors_consistent(tmp_path: Path) -> None:
    import multiprocessing as mp

    home = tmp_path / "concurrent-home"
    paths = Paths(home)
    save_config(paths, AppConfig(vector_backend="hashing", watcher_enabled=False, chunk_chars=400, overlap_chars=40))
    corpus = tmp_path / "concurrent-corpus"
    corpus.mkdir()
    # Multiple chunks make partial vector/catalog interleaving observable through count mismatches.
    (corpus / "race.md").write_text(
        "# Concurrent Writer Test\n\n" + ("atomic catalog vector invariant semaphore lighthouse\n" * 100),
        encoding="utf-8",
    )

    ctx = mp.get_context("spawn")
    gate = ctx.Barrier(2)
    out = ctx.Queue()
    procs = [ctx.Process(target=_concurrent_sync_worker, args=(str(home), str(corpus), gate, out)) for _ in range(2)]
    for proc in procs:
        proc.start()
    results = [out.get(timeout=30) for _ in procs]
    for proc in procs:
        proc.join(timeout=30)
        assert proc.exitcode == 0
    assert all("error" not in result and result["errors"] == 0 for result in results)

    rt = Runtime(home=home, start_watcher=False)
    try:
        status = rt.status()
        assert status["sources"] == 1
        assert status["chunks"] > 1
        assert status["vectors"] == status["chunks"]
        assert status["vector_count_matches_chunks"] is True
        hits = rt.search.hybrid_search("semaphore lighthouse", top_k=5)
        assert hits
    finally:
        rt.close()


def test_changed_unparseable_file_fails_closed_without_stale_search_results(tmp_path: Path) -> None:
    corpus = tmp_path / "fail-closed"
    corpus.mkdir()
    source = corpus / "notes.docx"
    doc = Document()
    doc.add_paragraph("last good stale sentinel zirconium theorem")
    doc.save(source)
    rt = make_runtime(tmp_path)
    try:
        rt.add_root(str(corpus), sync_now=True)
        assert any(Path(h.path) == source.resolve() for h in rt.search.hybrid_search("zirconium theorem", top_k=5))

        # A .docx is a ZIP container; these bytes are deliberately invalid real input.
        source.write_bytes(b"this is no longer a valid docx container")
        result = rt.indexer.sync_root(corpus)
        assert result.errors == 1
        row = rt.catalog.source_by_path(source)
        assert row is not None and row["status"] == "error"
        assert rt.status()["vectors"] == rt.status()["chunks"] == 0
        assert not rt.search.hybrid_search("zirconium theorem", top_k=5)
    finally:
        rt.close()


def test_deleted_directory_root_can_be_recreated_and_reindexed(tmp_path: Path) -> None:
    import shutil

    corpus = tmp_path / "restorable-root"
    corpus.mkdir()
    first = corpus / "first.md"
    first.write_text("root restoration first token", encoding="utf-8")
    rt = make_runtime(tmp_path)
    try:
        rt.add_root(str(corpus), sync_now=True)
        shutil.rmtree(corpus)
        rt.watcher._sync_directory_event(corpus)
        assert rt.status()["sources"] == 0
        assert str(corpus.resolve()) in rt.status()["roots"]

        corpus.mkdir()
        second = corpus / "second.md"
        second.write_text("root restoration second cobalt token", encoding="utf-8")
        rt.watcher._sync_directory_event(corpus)
        hits = rt.search.hybrid_search("cobalt token", top_k=5)
        assert any(Path(h.path) == second.resolve() for h in hits)
    finally:
        rt.close()


def test_offline_deleted_root_is_pruned_on_next_sync_but_remains_configured(tmp_path: Path) -> None:
    import shutil

    corpus = tmp_path / "offline-deleted-root"
    corpus.mkdir()
    note = corpus / "offline.md"
    note.write_text("offline stale sentinel vanadium paging", encoding="utf-8")
    rt = make_runtime(tmp_path)
    try:
        rt.add_root(str(corpus), sync_now=True)
        assert rt.status()["sources"] == 1
        shutil.rmtree(corpus)

        result = rt.indexer.sync_all()
        assert result["total"]["deleted"] == 1
        assert rt.status()["sources"] == 0
        assert rt.status()["chunks"] == 0
        assert rt.status()["vectors"] == 0
        assert str(corpus.resolve()) in rt.status()["roots"]
        assert not rt.search.hybrid_search("vanadium paging", top_k=5)

        corpus.mkdir()
        restored = corpus / "restored.md"
        restored.write_text("restored root tungsten attention", encoding="utf-8")
        recovered = rt.indexer.sync_all()
        assert recovered["total"]["indexed"] == 1
        assert any(Path(h.path) == restored.resolve() for h in rt.search.hybrid_search("tungsten attention", top_k=5))
    finally:
        rt.close()


def test_excluded_directory_names_are_case_insensitive(tmp_path: Path) -> None:
    corpus = tmp_path / "case-excludes"
    corpus.mkdir()
    (corpus / "visible.md").write_text("visible retrieval cobalt", encoding="utf-8")
    hidden = corpus / ".GIT"
    hidden.mkdir()
    (hidden / "secret.md").write_text("must never index hidden tungsten", encoding="utf-8")

    rt = make_runtime(tmp_path)
    try:
        rt.add_root(str(corpus), sync_now=True)
        assert rt.status()["sources"] == 1
        assert rt.search.hybrid_search("visible cobalt", top_k=5)
        assert rt.catalog.sources_under_root(corpus)[0]["path"] == str((corpus / "visible.md").resolve())
        lexical = rt.catalog.lexical_search("hidden tungsten", limit=5)
        assert lexical == []
    finally:
        rt.close()


def test_reconciliation_repairs_orphan_vectors_even_with_no_roots(tmp_path: Path) -> None:
    rt = make_runtime(tmp_path)
    try:
        # Simulate a crash that left one derivative vector without any catalog chunk/root.
        blob = (b"\x00\x00" * rt.vectors.dim)
        with rt.vectors._conn:
            rt.vectors._conn.execute("INSERT INTO vectors(key, vec) VALUES(?, ?)", (123456789, blob))
            rt.vectors._invalidate()
        assert rt.vectors.count() == 1
        assert rt.catalog.roots() == []
        assert rt.catalog.stats()["chunks"] == 0

        rt._safe_reconcile()
        assert rt.vectors.count() == 0
        assert rt.status()["vector_count_matches_chunks"] is True
    finally:
        rt.close()


def test_watch_command_indexes_new_files_in_background_process(tmp_path: Path) -> None:
    import pytest
    pytest.importorskip("watchdog")
    import os
    import subprocess
    import sys
    import time

    home = tmp_path / "watch-home"
    paths = Paths(home)
    save_config(paths, AppConfig(vector_backend="hashing", watcher_enabled=True, chunk_chars=500, overlap_chars=60))
    corpus = tmp_path / "watch-corpus"
    corpus.mkdir()

    setup = Runtime(home=home, start_watcher=False)
    try:
        setup.add_root(str(corpus), sync_now=True)
    finally:
        setup.close()

    env = os.environ.copy()
    env["PYTHONPATH"] = str((Path(__file__).resolve().parents[1] / "src").resolve())
    proc = subprocess.Popen(
        [sys.executable, "-m", "study_retriever.cli", "--home", str(home), "watch"],
        cwd=Path(__file__).resolve().parents[1],
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    try:
        ready_deadline = time.time() + 8.0
        ready = False
        assert proc.stdout is not None
        while time.time() < ready_deadline:
            line = proc.stdout.readline()
            if not line and proc.poll() is not None:
                break
            if '"watching": true' in line:
                ready = True
                break
        assert ready, (proc.stderr.read() if proc.poll() is not None and proc.stderr else "watch process did not become ready")

        note = corpus / "live.md"
        note.write_text("background watcher indexes indigo transformer sentinel", encoding="utf-8")

        deadline = time.time() + 8.0
        found = False
        while time.time() < deadline:
            reader = Runtime(home=home, start_watcher=False)
            try:
                found = any(
                    Path(hit.path) == note.resolve()
                    for hit in reader.search.hybrid_search("indigo transformer sentinel", top_k=5)
                )
            finally:
                reader.close()
            if found:
                break
            time.sleep(0.25)
        assert found, (proc.stderr.read() if proc.poll() is not None and proc.stderr else "watch process still running")
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=5)
        except subprocess.TimeoutExpired:
            proc.kill()
            proc.wait(timeout=5)


def test_typst_quarto_and_image_ocr_support(tmp_path) -> None:
    import shutil

    from study_retriever.parsers import _legible, is_supported, parse_document

    typ = tmp_path / "notes.typ"
    typ.write_text("= Gradient descent\nThe learning rate controls step size.\n", encoding="utf-8")
    qmd = tmp_path / "lecture.qmd"
    qmd.write_text("---\ntitle: Backprop\n---\n\n# Chain rule\nGradients flow backwards.\n", encoding="utf-8")
    for path, needle in ((typ, "learning rate"), (qmd, "Gradients flow")):
        assert is_supported(path)
        assert needle in parse_document(path).blocks[0].text
    assert is_supported(tmp_path / "x.png") is False  # missing file is never supported
    assert _legible("Gradient descent follows the negative gradient downhill each step")
    assert not _legible("eo a aia ie cia ae 4 eeee ae | e So")
    if shutil.which("tesseract") is None:
        return
    import fitz

    png = tmp_path / "caption.png"
    doc = fitz.open()
    page = doc.new_page(width=600, height=200)
    page.insert_text((20, 100), "Gradient descent follows the negative gradient downhill", fontsize=20)
    page.get_pixmap(dpi=150).save(png)
    parsed = parse_document(png)
    assert "gradient" in parsed.blocks[0].text.lower()


def test_svg_npz_and_markdown_image_cleanup(tmp_path) -> None:
    import numpy as np

    from study_retriever.parsers import parse_document

    svg = tmp_path / "fig.svg"
    svg.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg"><title>Loss surface</title>'
        '<text x="1">Swish <tspan>activation</tspan></text><path d="M0 0"/></svg>',
        encoding="utf-8",
    )
    assert "Swish activation" in parse_document(svg).blocks[0].text
    empty = tmp_path / "paths.svg"
    empty.write_text('<svg xmlns="http://www.w3.org/2000/svg"><path d="M0 0"/></svg>', encoding="utf-8")
    assert parse_document(empty).blocks == []

    npz = tmp_path / "data.npz"
    np.savez(npz, features=np.zeros((4, 8), dtype="float32"), breeds=np.array(["beagle", "pug"]))
    text = parse_document(npz).blocks[0].text
    assert "features: shape (4, 8)" in text and "beagle" in text

    md = tmp_path / "book.md"
    md.write_text("Intro\n\n![Page 7 image](images/p7.png)\n\n![Dirichlet samples](images/d.png)\n\nOutro\n", encoding="utf-8")
    body = parse_document(md).blocks[0].text
    assert "images/" not in body and "Dirichlet samples" in body and "Page 7" not in body
